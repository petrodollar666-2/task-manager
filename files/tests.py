import io
from datetime import datetime, timezone

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.urls import NoReverseMatch, reverse

from approvals.models import Proposal
from files.extract import MAX_CHARS, extract_text
from files.management.commands.sync_drive import DriveSyncer
from files.models import Document
from files.search import attach_snippets, fold, highlight, search_documents
from tasks.models import Project

User = get_user_model()


def tasks_detail_available():
    try:
        reverse("tasks:detail", args=[1])
        return True
    except NoReverseMatch:
        return False


needs_tasks_detail = pytest.mark.skipif(
    not tasks_detail_available(), reason="URL tasks:detail chưa có (app tasks đang được xây dựng)"
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture
def boss(db):
    return User.objects.create_user("sep", password="pw", full_name="Trưởng khoa", role="admin")


@pytest.fixture
def alice(db):
    return User.objects.create_user("alice", password="pw", full_name="Nguyễn Thị An", role="member")


@pytest.fixture
def bob(db):
    return User.objects.create_user("bob", password="pw", full_name="Trần Văn Bình", role="member")


@pytest.fixture(autouse=True)
def media_tmp(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "media"
    return settings.MEDIA_ROOT


def make_doc(owner=None, **kwargs):
    kwargs.setdefault("name", "Tài liệu")
    return Document.objects.create(owner=owner, **kwargs)


def search_names(q, qs=None):
    results, _ = search_documents(qs if qs is not None else Document.objects.all(), q)
    return [d.name for d in results]


# ---------------------------------------------------------------------------
# Trích xuất nội dung
# ---------------------------------------------------------------------------
def test_extract_docx_with_tables(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Kế hoạch đào tạo năm 2026")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Học phần"
    table.cell(0, 1).text = "Kinh tế lượng"
    table.cell(1, 0).text = "Giảng viên"
    table.cell(1, 1).text = "Nguyễn Văn A"
    path = tmp_path / "ke-hoach.docx"
    d.save(path)

    text = extract_text(path, "ke-hoach.docx")
    assert "Kế hoạch đào tạo năm 2026" in text
    assert "Kinh tế lượng" in text
    assert "Nguyễn Văn A" in text

    # Cũng đọc được từ file object
    with open(path, "rb") as fh:
        assert "Kinh tế lượng" in extract_text(fh, "ke-hoach.docx")


def test_extract_xlsx_all_sheets(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Dự toán"
    ws.append(["Khoản mục", "Số tiền"])
    ws.append(["Thuê hội trường", 15000000])
    ws2 = wb.create_sheet("Ghi chú")
    ws2["B3"] = "Quyết toán hội thảo"
    path = tmp_path / "du-toan.xlsx"
    wb.save(path)

    text = extract_text(str(path), "du-toan.xlsx")
    assert "Thuê hội trường" in text
    assert "15000000" in text
    assert "Quyết toán hội thảo" in text
    assert "Dự toán" in text


def test_extract_txt_csv_md(tmp_path):
    p = tmp_path / "ghi-chu.txt"
    p.write_text("Biên bản họp bộ môn\nNội dung: đào tạo", encoding="utf-8")
    assert "Biên bản họp bộ môn" in extract_text(p, "ghi-chu.txt")
    assert "a,b" in extract_text(io.BytesIO(b"a,b\n1,2"), "x.csv")
    assert "Tiêu đề" in extract_text(io.BytesIO("# Tiêu đề".encode()), "README.md")
    # Đoán loại theo mime khi không có đuôi file
    assert "xin chào" in extract_text(io.BytesIO("xin chào".encode()), "noext", "text/plain")


def test_extract_never_crashes_and_caps():
    assert extract_text(io.BytesIO(b"not a pdf"), "hong.pdf") == ""
    assert extract_text(io.BytesIO(b"PK garbage"), "hong.docx") == ""
    assert extract_text(io.BytesIO(b"\x00\x01"), "anh.png") == ""
    assert extract_text("/khong/ton/tai.txt", "tai.txt") == ""
    big = ("chữ " * 400_000).encode()
    text = extract_text(io.BytesIO(b"a\x00b " + big), "lon.txt")
    assert len(text) <= MAX_CHARS
    assert "\x00" not in text


def test_extract_pdf(tmp_path):
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    path = tmp_path / "trang.pdf"
    with open(path, "wb") as fh:
        w.write(fh)
    assert extract_text(path, "trang.pdf") == ""  # trang trắng: không lỗi


# ---------------------------------------------------------------------------
# Tìm kiếm
# ---------------------------------------------------------------------------
@pytest.mark.django_db
def test_search_accent_insensitive_both_directions(alice):
    make_doc(alice, name="Kế hoạch tài chính 2026")
    make_doc(alice, name="ke hoach khong dau")
    make_doc(alice, name="Lịch thi học kỳ")

    names = search_names("ke hoach tai chinh")
    assert names == ["Kế hoạch tài chính 2026"]

    names = search_names("Kế hoạch")
    assert set(names) == {"Kế hoạch tài chính 2026", "ke hoach khong dau"}

    assert search_names("KẾ HOẠCH TÀI CHÍNH") == ["Kế hoạch tài chính 2026"]


@pytest.mark.django_db
def test_search_d_stroke(alice):
    make_doc(alice, name="Đào tạo sau đại học")
    make_doc(alice, name="Dao động điều hoà")
    assert "Đào tạo sau đại học" in search_names("dao tao")
    assert search_names("dao tao") == ["Đào tạo sau đại học"]
    assert search_names("đào tạo") == ["Đào tạo sau đại học"]


@pytest.mark.django_db
def test_search_in_content_and_description(alice):
    make_doc(alice, name="bao-cao.pdf", extracted_text="Báo cáo tổng kết hội thảo quốc tế về kinh tế số")
    make_doc(alice, name="khac.docx", description="Tờ trình mua sắm thiết bị phòng máy")
    make_doc(alice, name="rac.txt", extracted_text="không liên quan")

    assert search_names("hoi thao quoc te") == ["bao-cao.pdf"]
    assert search_names("thiet bi") == ["khac.docx"]


@pytest.mark.django_db
def test_search_prefix_and_ranking(alice):
    make_doc(alice, name="Kế hoạch tài chính")
    make_doc(alice, name="Ghi chú", extracted_text="có nhắc đến tài chính ở nội dung")
    names = search_names("tai chi")
    assert set(names) == {"Kế hoạch tài chính", "Ghi chú"}
    assert names[0] == "Kế hoạch tài chính"  # khớp ở tên (trọng số A) xếp trước
    assert search_names("hoa") == ["Kế hoạch tài chính"]


@pytest.mark.django_db
def test_search_trigram_fallback_and_empty(alice):
    make_doc(alice, name="Kế hoạch tài chính")
    assert search_names("ke hoach tai chinhh") == ["Kế hoạch tài chính"]
    assert search_names("xyzxyz") == []
    assert search_names("!!!") == []
    assert search_names("") == ["Kế hoạch tài chính"]


@pytest.mark.django_db
def test_search_websearch_syntax(alice):
    make_doc(alice, name="Kế hoạch đào tạo")
    make_doc(alice, name="Kế hoạch tài chính")
    assert search_names("ke hoach -tai") == ["Kế hoạch đào tạo"]
    assert search_names('"hoach dao"') == ["Kế hoạch đào tạo"]


@pytest.mark.django_db
def test_search_vector_updates_on_edit(alice):
    doc = make_doc(alice, name="Cũ")
    assert search_names("moi") == []
    doc.name = "Mới"
    doc.save()
    assert search_names("moi") == ["Mới"]


@pytest.mark.django_db
def test_snippet_highlights_original_text_and_escapes(alice):
    text = ("Mở đầu. " * 40) + "Hội thảo <script>alert(1)</script> về Đào tạo" + (" kết thúc." * 40)
    make_doc(alice, name="Đào tạo", extracted_text=text)
    results, terms = search_documents(Document.objects.all(), "dao tao")
    [doc] = attach_snippets(results, terms)
    html = str(doc.snippet_html)
    assert '<mark class="rounded bg-yellow-200 px-0.5">Đào</mark>' in html
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert html.startswith("…")
    assert "<mark" in str(doc.name_html)


def test_fold_and_highlight_helpers():
    assert fold("Đào Tạo Ưu Tiên") == "dao tao uu tien"
    assert len(fold("Kế hoạch")) == len("Kế hoạch")
    assert str(highlight("a<b> Kế hoạch", ["ke"])) == 'a&lt;b&gt; <mark class="rounded bg-yellow-200 px-0.5">Kế</mark> hoạch'


# ---------------------------------------------------------------------------
# Quyền truy cập
# ---------------------------------------------------------------------------
@pytest.mark.django_db
def test_visibility_rules(boss, alice, bob):
    own = make_doc(alice, name="Của Alice")
    project = Project.objects.create(title="Workshop", owner=bob)
    proposal = Proposal.objects.create(title="Tờ trình", owner=bob)
    on_project = make_doc(alice, name="Gắn đầu việc của Bob", project=project)
    on_proposal = make_doc(alice, name="Gắn tờ trình của Bob", proposal=proposal)
    orphan = make_doc(None, name="Từ Drive", source="drive")

    assert set(Document.objects.visible_to(boss)) == {own, on_project, on_proposal, orphan}
    assert set(Document.objects.visible_to(alice)) == {own, on_project, on_proposal}
    assert set(Document.objects.visible_to(bob)) == {on_project, on_proposal}
    assert on_project.can_edit(alice) and on_project.can_edit(boss)
    assert not on_project.can_edit(bob)


@pytest.mark.django_db
def test_member_cannot_see_or_download_others(client, alice, bob):
    doc = make_doc(alice, name="Bí mật",
                   file=SimpleUploadedFile("bi-mat.txt", b"noi dung"), original_filename="bi-mat.txt")
    client.force_login(bob)
    resp = client.get(reverse("files:search"), {"q": "bi mat"})
    assert resp.status_code == 200
    assert "Bí mật" not in resp.content.decode()
    assert "Tìm thấy <strong class=\"text-slate-900\">0</strong>" in resp.content.decode()
    assert client.get(reverse("files:detail", args=[doc.pk])).status_code == 404
    assert client.get(reverse("files:download", args=[doc.pk])).status_code == 404
    assert client.get(reverse("files:edit", args=[doc.pk])).status_code == 404
    assert client.post(reverse("files:delete", args=[doc.pk])).status_code == 404
    assert Document.objects.filter(pk=doc.pk).exists()


@pytest.mark.django_db
def test_boss_can_search_and_download(client, boss, alice):
    doc = make_doc(alice, name="Kế hoạch tài chính",
                   file=SimpleUploadedFile("kh.txt", "Nội dung kế hoạch".encode()),
                   original_filename="Kế hoạch tài chính.txt", mime_type="text/plain")
    client.force_login(boss)
    resp = client.get(reverse("files:search"), {"q": "ke hoach"})
    body = resp.content.decode()
    assert resp.status_code == 200
    assert "Tìm thấy <strong class=\"text-slate-900\">1</strong>" in body
    assert reverse("files:download", args=[doc.pk]) in body
    assert "<mark" in body

    resp = client.get(reverse("files:download", args=[doc.pk]))
    assert resp.status_code == 200
    assert "attachment" in resp["Content-Disposition"]
    assert "K%E1%BA%BF%20ho%E1%BA%A1ch" in resp["Content-Disposition"]
    assert b"".join(resp.streaming_content) == "Nội dung kế hoạch".encode()

    assert client.get(reverse("files:detail", args=[doc.pk])).status_code == 200
    assert client.get(reverse("files:edit", args=[doc.pk])).status_code == 200


@pytest.mark.django_db
def test_project_owner_can_view_but_not_edit(client, alice, bob):
    project = Project.objects.create(title="Workshop", owner=bob)
    doc = make_doc(alice, name="Chương trình", project=project)
    client.force_login(bob)
    if tasks_detail_available():
        assert client.get(reverse("files:detail", args=[doc.pk])).status_code == 200
    assert client.get(reverse("files:edit", args=[doc.pk])).status_code == 403
    assert client.post(reverse("files:delete", args=[doc.pk])).status_code == 403


@pytest.mark.django_db
def test_login_required(client):
    resp = client.get(reverse("files:search"))
    assert resp.status_code == 302
    assert reverse("login") in resp["Location"]


# ---------------------------------------------------------------------------
# Tải lên / sửa / xoá
# ---------------------------------------------------------------------------
@pytest.mark.django_db
def test_upload_creates_document_with_extracted_text(client, alice, bob, media_tmp):
    import docx

    d = docx.Document()
    d.add_paragraph("Đề cương học phần Kinh tế vĩ mô")
    buf = io.BytesIO()
    d.save(buf)
    upload = SimpleUploadedFile(
        "Đề cương.docx", buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    client.force_login(alice)
    resp = client.post(reverse("files:upload"), {"file": upload, "name": "", "description": "Học kỳ 1"})
    assert resp.status_code == 302, resp.content.decode()

    doc = Document.objects.get()
    assert doc.owner == alice
    assert doc.name == "Đề cương.docx"
    assert doc.original_filename == "Đề cương.docx"
    assert doc.source == "upload"
    assert "Kinh tế vĩ mô" in doc.extracted_text
    assert doc.file.path.startswith(str(media_tmp))
    assert search_names("kinh te vi mo") == ["Đề cương.docx"]
    assert resp["Location"] == reverse("files:detail", args=[doc.pk])


@pytest.mark.django_db
def test_upload_limits_projects_to_own(client, alice, bob):
    mine = Project.objects.create(title="Của Alice", owner=alice)
    other = Project.objects.create(title="Của Bob", owner=bob)
    client.force_login(alice)
    body = client.get(reverse("files:upload")).content.decode()
    assert "Của Alice" in body and "Của Bob" not in body

    resp = client.post(reverse("files:upload"), {
        "file": SimpleUploadedFile("a.txt", b"x"), "name": "A", "project": other.pk,
    })
    assert resp.status_code == 200  # lựa chọn không hợp lệ
    assert not Document.objects.exists()

    resp = client.post(reverse("files:upload"), {
        "file": SimpleUploadedFile("a.txt", b"x"), "name": "A", "project": mine.pk,
    })
    assert resp.status_code == 302
    assert Document.objects.get().project == mine


@pytest.mark.django_db
def test_edit_replace_file_and_delete(client, alice):
    client.force_login(alice)
    client.post(reverse("files:upload"), {"file": SimpleUploadedFile("cu.txt", "nội dung cũ".encode())})
    doc = Document.objects.get()
    old_path = doc.file.path

    resp = client.post(reverse("files:edit", args=[doc.pk]), {
        "name": "Tên mới", "description": "",
        "file": SimpleUploadedFile("moi.txt", "nội dung mới về ngân sách".encode()),
    })
    assert resp.status_code == 302
    doc.refresh_from_db()
    assert doc.name == "Tên mới"
    assert doc.original_filename == "moi.txt"
    assert "ngân sách" in doc.extracted_text
    assert search_names("ngan sach") == ["Tên mới"]
    import os
    assert not os.path.exists(old_path)

    assert client.get(reverse("files:delete", args=[doc.pk])).status_code == 200
    resp = client.post(reverse("files:delete", args=[doc.pk]))
    assert resp.status_code == 302
    assert not Document.objects.exists()
    assert not os.path.exists(doc.file.path)


@needs_tasks_detail
@pytest.mark.django_db
def test_search_page_renders_project_link(client, boss):
    project = Project.objects.create(title="Workshop FTU", owner=boss)
    make_doc(boss, name="Chương trình workshop", project=project)
    client.force_login(boss)
    body = client.get(reverse("files:search")).content.decode()
    assert reverse("tasks:detail", args=[project.pk]) in body


# ---------------------------------------------------------------------------
# Lệnh quản trị
# ---------------------------------------------------------------------------
@pytest.mark.django_db
def test_sync_drive_not_configured(settings, capsys):
    settings.GOOGLE_SERVICE_ACCOUNT_FILE = ""
    settings.GOOGLE_DRIVE_FOLDER_ID = ""
    out = io.StringIO()
    call_command("sync_drive", stdout=out)  # không ném lỗi / SystemExit
    text = out.getvalue()
    assert "Chưa cấu hình" in text
    assert "GOOGLE_SERVICE_ACCOUNT_FILE" in text
    assert "GOOGLE_DRIVE_FOLDER_ID" in text


class _Req:
    def __init__(self, value):
        self.value = value

    def execute(self):
        return self.value


class FakeDrive:
    """Giả lập tối thiểu service.files() của Drive API v3."""

    def __init__(self, tree, contents):
        self.tree, self.contents, self.downloads = tree, contents, []

    def files(self):
        return self

    def list(self, q, pageToken=None, **kw):
        folder = q.split("'")[1]
        return _Req({"files": self.tree.get(folder, [])})

    def get_media(self, fileId, **kw):
        self.downloads.append(fileId)
        return _Req(self.contents[fileId])

    def export_media(self, fileId, mimeType):
        self.downloads.append(fileId)
        return _Req(self.contents[fileId])


@pytest.mark.django_db
def test_drive_syncer_upserts_recursively():
    t1 = "2026-01-01T10:00:00.000Z"
    tree = {
        "root": [
            {"id": "f1", "name": "Ghi chú.txt", "mimeType": "text/plain", "modifiedTime": t1,
             "webViewLink": "https://drive.google.com/file/d/f1/view", "size": "20"},
            {"id": "sub", "name": "Con", "mimeType": "application/vnd.google-apps.folder"},
        ],
        "sub": [
            {"id": "g1", "name": "Kế hoạch đào tạo", "mimeType": "application/vnd.google-apps.document",
             "modifiedTime": t1, "webViewLink": "https://docs.google.com/document/d/g1/edit"},
        ],
    }
    contents = {"f1": "Biên bản họp".encode(), "g1": "Nội dung kế hoạch đào tạo".encode()}
    drive = FakeDrive(tree, contents)

    stats = DriveSyncer(drive, log=lambda *a: None).sync("root")
    assert stats["created"] == 2 and stats["errors"] == 0
    g1 = Document.objects.get(drive_file_id="g1")
    assert g1.source == "drive"
    assert g1.drive_modified_time == datetime(2026, 1, 1, 10, tzinfo=timezone.utc)
    assert "Nội dung kế hoạch" in g1.extracted_text
    assert search_names("bien ban") == ["Ghi chú.txt"]

    # Chạy lại: không đổi → bỏ qua, không tải lại
    drive.downloads.clear()
    stats = DriveSyncer(drive, log=lambda *a: None).sync("root")
    assert stats["skipped"] == 2 and drive.downloads == []

    # File thay đổi → cập nhật
    tree["root"][0]["modifiedTime"] = "2026-02-01T10:00:00Z"
    contents["f1"] = "Biên bản mới về ngân sách".encode()
    stats = DriveSyncer(drive, log=lambda *a: None).sync("root")
    assert stats["updated"] == 1 and stats["skipped"] == 1
    assert Document.objects.count() == 2
    assert search_names("ngan sach") == ["Ghi chú.txt"]


@pytest.mark.django_db
def test_reindex_documents(alice):
    doc = make_doc(alice, name="x", file=SimpleUploadedFile("x.txt", "học bổng khuyến khích".encode()),
                   original_filename="x.txt")
    assert doc.extracted_text == ""
    Document.objects.update(search_vector=None)
    out = io.StringIO()
    call_command("reindex_documents", stdout=out)
    doc.refresh_from_db()
    assert "học bổng" in doc.extracted_text
    assert search_names("hoc bong") == ["x"]
    assert "Xong" in out.getvalue()
