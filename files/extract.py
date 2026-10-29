"""Trích xuất văn bản từ file để phục vụ tìm kiếm toàn văn.

    extract_text(path_or_fileobj, filename, mime=None) -> str

Hỗ trợ .pdf (pypdf), .docx (python-docx, gồm cả bảng), .xlsx (openpyxl, mọi
sheet), .txt/.csv/.md/.tsv. Không bao giờ ném lỗi: thất bại thì trả về "".
"""

import io
import logging
import os

logger = logging.getLogger(__name__)

MAX_CHARS = 1_000_000

TEXT_EXTS = {"txt", "csv", "md", "tsv", "text", "markdown"}
MIME_TO_EXT = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel.sheet.macroenabled.12": "xlsm",
    "text/plain": "txt",
    "text/csv": "csv",
    "text/markdown": "md",
    "text/tab-separated-values": "tsv",
}
SUPPORTED_EXTS = {"pdf", "docx", "xlsx", "xlsm"} | TEXT_EXTS


class _Collector:
    """Gom các đoạn văn bản, dừng khi đạt giới hạn ký tự."""

    def __init__(self, limit=MAX_CHARS):
        self.parts = []
        self.size = 0
        self.limit = limit

    @property
    def full(self):
        return self.size >= self.limit

    def add(self, text):
        if not text or self.full:
            return
        text = str(text)
        self.parts.append(text)
        self.size += len(text) + 1

    def result(self):
        return "\n".join(self.parts)


def detect_kind(filename="", mime=None):
    ext = os.path.splitext(filename or "")[1].lower().lstrip(".")
    if ext in SUPPORTED_EXTS:
        return ext
    if mime:
        mime = mime.split(";")[0].strip().lower()
        if mime in MIME_TO_EXT:
            return MIME_TO_EXT[mime]
        if mime.startswith("text/"):
            return "txt"
    return ext


def is_supported(filename="", mime=None):
    return detect_kind(filename, mime) in SUPPORTED_EXTS


def _read_bytes(source):
    if isinstance(source, (str, os.PathLike)):
        with open(source, "rb") as fh:
            return fh.read()
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if hasattr(source, "seek"):
        try:
            source.seek(0)
        except (OSError, ValueError):
            pass
    if hasattr(source, "chunks"):  # Django UploadedFile / File
        return b"".join(source.chunks())
    return source.read()


def _pdf(data, out):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:  # noqa: BLE001
            return
    for page in reader.pages:
        if out.full:
            break
        try:
            out.add(page.extract_text())
        except Exception:  # noqa: BLE001 – một trang lỗi không làm hỏng cả file
            logger.debug("Không đọc được một trang PDF", exc_info=True)


def _docx(data, out):
    import docx

    document = docx.Document(io.BytesIO(data))
    for para in document.paragraphs:
        out.add(para.text)
    for table in document.tables:
        _docx_table(table, out)
    for section in document.sections:
        for part in (section.header, section.footer):
            try:
                for para in part.paragraphs:
                    out.add(para.text)
            except Exception:  # noqa: BLE001
                pass


def _docx_table(table, out):
    for row in table.rows:
        if out.full:
            return
        seen = set()
        cells = []
        for cell in row.cells:
            # Ô gộp (merged) xuất hiện lặp lại → chỉ lấy một lần.
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append(cell.text.strip())
            for nested in cell.tables:
                _docx_table(nested, out)
        out.add("\t".join(c for c in cells if c))


def _xlsx(data, out):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if out.full:
                break
            out.add(ws.title)
            for row in ws.iter_rows(values_only=True):
                if out.full:
                    break
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if values:
                    out.add("\t".join(values))
    finally:
        wb.close()


def _text(data, out):
    for encoding in ("utf-8-sig", "utf-16", "cp1258"):
        try:
            if encoding == "utf-16" and not data[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            out.add(data.decode(encoding))
            return
        except (UnicodeDecodeError, LookupError):
            continue
    out.add(data.decode("utf-8", errors="replace"))


HANDLERS = {
    "pdf": _pdf,
    "docx": _docx,
    "xlsx": _xlsx,
    "xlsm": _xlsx,
    **{ext: _text for ext in TEXT_EXTS},
}


def clean_text(text, limit=MAX_CHARS):
    # Postgres không lưu được ký tự NUL.
    text = text.replace("\x00", "")
    return text[:limit].strip()


def extract_text(path_or_fileobj, filename="", mime=None):
    """Trả về nội dung văn bản của file (tối đa MAX_CHARS ký tự), hoặc "" nếu
    không hỗ trợ / lỗi."""
    if not filename and isinstance(path_or_fileobj, (str, os.PathLike)):
        filename = os.fspath(path_or_fileobj)
    if not filename:
        filename = getattr(path_or_fileobj, "name", "") or ""
    kind = detect_kind(filename, mime)
    handler = HANDLERS.get(kind)
    if handler is None:
        return ""
    try:
        data = _read_bytes(path_or_fileobj)
        out = _Collector()
        handler(data, out)
        return clean_text(out.result())
    except Exception:  # noqa: BLE001 – không bao giờ làm hỏng luồng tải lên
        logger.warning("Không trích xuất được nội dung từ %s", filename, exc_info=True)
        return ""
